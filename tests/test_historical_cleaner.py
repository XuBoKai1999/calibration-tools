import csv
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import historical_cleaner as cleaner  # noqa: E402


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


class HistoricalCleanerTests(unittest.TestCase):
    def raw(self, system, name):
        return ROOT / "past" / system / "source" / "raw_data" / name

    def convert(self, system, name, output):
        return cleaner.convert_one(ROOT, system, self.raw(system, name), Path(output))

    def test_all_representative_sources_are_unchanged_and_manifest_hashes_match(self):
        sources = []
        for system, names in cleaner.REPRESENTATIVES.items():
            for name in names:
                raw = self.raw(system, name)
                sources.append(raw)
                sources.extend(cleaner.find_reports(ROOT / "past" / system, raw.stem))
        before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in sources}
        with tempfile.TemporaryDirectory() as directory:
            cleaner.run_representatives(ROOT, Path(directory))
            for system, names in cleaner.REPRESENTATIVES.items():
                for name in names:
                    manifest = json.loads((Path(directory) / system / Path(name).stem / "manifest.json").read_text(encoding="utf-8"))
                    for item in manifest["sources"]:
                        source = ROOT / "past" / system / item["path"]
                        self.assertEqual(before[source], item["sha256"])
        self.assertEqual(before, {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in sources})

    def test_package_names_missing_report_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self.convert("E05", "E170381.xls", directory)
            package = Path(directory) / "E05" / "E170381"
            self.assertEqual({"manifest.json", "raw.csv", "context.csv"}, {p.name for p in package.iterdir()})
            manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
            self.assertIsNone(manifest["artifacts"]["report"])
            self.assertEqual("FAIL_AMBIGUOUS", result["conversion_status"])
            self.assertEqual("known_raw_report_mismatch", manifest["conflicts"][0]["concept"])
            with self.assertRaises(FileExistsError):
                self.convert("E05", "E170381.xls", directory)

    def test_unknown_extra_sheet_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "unknown.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active["A1"] = "unexpected"
            workbook.create_sheet("additional")["A1"] = "significant"
            workbook.save(path)
            with self.assertRaises(cleaner.UnsupportedStructure):
                cleaner.detect_generation("E05", path, cleaner.load_book(path))

    def test_e05_observations_and_auxiliary_meter(self):
        with tempfile.TemporaryDirectory() as directory:
            self.convert("E05", "E250527.xls", directory)
            package = Path(directory) / "E05" / "E250527"
            raw = read_csv(package / "raw.csv")
            context = read_csv(package / "context.csv")
            self.assertEqual(60, len(raw))
            self.assertTrue(all(row["mode"] == "direct_voltage" for row in raw))
            meter = {row["key"]: row["value"] for row in context if row["scope_id"] == "meter_01"}
            self.assertEqual({"manufacturer": "FLUKE", "model": "187", "serial_number": "94150148",
                              "instrument_role": "customer_meter"}, meter)
            self.assertNotIn("resolution", meter)
            self.assertNotIn("range", meter)

    def test_e05_g4_metadata_dates_and_selected_standards(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self.convert("E05", "E260384.xlsx", directory)
            package = Path(directory) / "E05" / "E260384"
            context = read_csv(package / "context.csv")
            case = {row["key"]: row["value"] for row in context if row["scope_type"] == "case"}
            self.assertEqual("交直流高壓分壓器", case["dut_name"])
            self.assertEqual("JRL", case["dut_manufacturer"])
            self.assertEqual("HVA-50", case["dut_model"])
            self.assertEqual("505", case["dut_serial_number"])
            instruments = {(row["scope_id"], row["key"]): row["value"] for row in context if row["scope_type"] == "instrument"}
            self.assertEqual("14872", instruments[("standard_divider_01", "serial_number")])
            self.assertEqual("2823A08446", instruments[("standard_voltmeter_01", "serial_number")])
            manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual("FAIL_AMBIGUOUS", result["conversion_status"])
            self.assertEqual("calibration_dates", manifest["conflicts"][0]["concept"])
            self.assertEqual(["2025-07-13", "2025-07-14", "2025-07-15", "2025-07-16"], manifest["conflicts"][0]["raw"])
            self.assertEqual(["2026-07-13", "2026-07-16"], manifest["conflicts"][0]["report"])

    def test_e05_report_fallback_and_conflict(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            raw_path = directory / "sample.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet["A5"], sheet["B5"] = "id", "dut"
            sheet["A7"], sheet["B7"] = "header", "2026-01-01"
            sheet["A8"], sheet["B8"] = "10 kV", 1.0
            sheet["A15"] = "23.0 1.5 45 10"
            sheet["P63"] = "helper"
            workbook.save(raw_path)
            report_path = directory / "report.docx"
            document = Document()
            document.add_paragraph("attached meter FLUKE/187/REPORT-SN")
            document.save(report_path)
            book = cleaner.load_book(raw_path)
            _, context, conflicts, _ = cleaner.clean_e05(raw_path, book, "E05-G4", report_path)
            meter = {row["key"]: row["value"] for row in context if row["scope_id"] == "meter_01"}
            self.assertEqual("REPORT-SN", meter["serial_number"])
            self.assertFalse(conflicts)
            workbook = openpyxl.load_workbook(raw_path)
            workbook.active["A20"] = "*FLUKE/187/RAW-SN"
            workbook.save(raw_path)
            _, _, conflicts, _ = cleaner.clean_e05(raw_path, cleaner.load_book(raw_path), "E05-G4", report_path)
            self.assertEqual("customer_meter_identity", conflicts[0]["concept"])

    def test_e07_direct_mode_and_no_default_frequency(self):
        with tempfile.TemporaryDirectory() as directory:
            self.convert("E07", "E200397.xls", directory)
            direct = Path(directory) / "E07" / "E200397"
            self.assertTrue(all(row["mode"] == "ac_high_voltage" for row in read_csv(direct / "raw.csv")))
            self.assertEqual(["60"], [row["value"] for row in read_csv(direct / "context.csv") if row["key"] == "frequency"])
        with tempfile.TemporaryDirectory() as directory:
            self.convert("E07", "E150076.xls", directory)
            transformer = Path(directory) / "E07" / "E150076"
            self.assertTrue(all(not row["frequency_value"] for row in read_csv(transformer / "raw.csv")))
            self.assertFalse(any(row["key"] == "frequency" for row in read_csv(transformer / "context.csv")))

    def test_e07_environment_standards_and_g4_date_recovery(self):
        with tempfile.TemporaryDirectory() as directory:
            self.convert("E07", "E250585.xls", directory)
            context = read_csv(Path(directory) / "E07" / "E250585" / "context.csv")
            self.assertEqual(["45"], [row["value"] for row in context if row["key"] == "humidity"])
            instruments = {(row["scope_id"], row["key"]): row["value"] for row in context if row["scope_type"] == "instrument"}
            self.assertEqual("143842", instruments[("standard_transformer_01", "serial_number")])
            self.assertEqual("5355303", instruments[("standard_ac_voltmeter_01", "serial_number")])
        with tempfile.TemporaryDirectory() as directory:
            result = self.convert("E07", "E260384.xlsx", directory)
            package = Path(directory) / "E07" / "E260384"
            self.assertEqual({"2026-07-15", "2026-07-16", "2026-07-17", "2026-07-20"},
                             {row["date"] for row in read_csv(package / "raw.csv")})
            self.assertEqual("PASS_WITH_UNRESOLVED", result["conversion_status"])
            manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
            self.assertTrue(any(item["concept"] == "unmapped_labelled_metadata" for item in manifest["unresolved"]))

    def test_e07_burden_scopes_and_setup_links(self):
        with tempfile.TemporaryDirectory() as directory:
            self.convert("E07", "E150076.xls", directory)
            package = Path(directory) / "E07" / "E150076"
            self.assertEqual(["15"], [row["value"] for row in read_csv(package / "context.csv")
                                      if row["scope_type"] == "run" and row["key"] == "burden"])
        with tempfile.TemporaryDirectory() as directory:
            result = self.convert("E07", "09807C02014.xls", directory)
            package = Path(directory) / "E07" / "09807C02014"
            context = read_csv(package / "context.csv")
            burdens = {(row["scope_id"], row["value"]) for row in context if row["key"] == "burden"}
            self.assertEqual({("setup_01", "1.25"), ("setup_02", "2.5")}, burdens)
            self.assertEqual({"setup_01", "setup_02"}, {row["setup_id"] for row in read_csv(package / "raw.csv")})
            self.assertEqual("FAIL_AMBIGUOUS", result["conversion_status"])

    def test_e27_observations_and_irreducible_summary(self):
        with tempfile.TemporaryDirectory() as directory:
            self.convert("E27", "E260073.xlsx", directory)
            rows = read_csv(Path(directory) / "E27" / "E260073" / "raw.csv")
            self.assertEqual({"dut", "reference_wafer"}, {row["sample_role"] for row in rows})
            self.assertTrue(all(row["source_stddev"] for row in rows))
            self.assertTrue(all(row["source_stddev_unit"] == "1" for row in rows))
            self.assertTrue(all("source-provided" in row["source_summary_note"] for row in rows))
            context = read_csv(Path(directory) / "E27" / "E260073" / "context.csv")
            values = {(row["scope_id"], row["key"]): row["value"] for row in context}
            self.assertEqual("US36001721", values[("standard_meter_01", "serial_number")])
            self.assertEqual("E260042", values[("standard_meter_01", "traceability_id")])
            self.assertEqual("62458", values[("standard_resistor_05", "serial_number")])
            self.assertEqual("E250675", values[("standard_resistor_05", "traceability_id")])
            self.assertEqual("1596.7", values[("reference_wafer_01", "certificate_value")])
            self.assertEqual("1.8", values[("reference_wafer_01", "certificate_uncertainty")])

    def test_e27_multisheet_is_not_merged(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self.convert("E27", "10407C02798-100.xls", directory)
            package = Path(directory) / "E27" / "10407C02798-100"
            manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual("FAIL_AMBIGUOUS", result["conversion_status"])
            self.assertIsNone(manifest["artifacts"]["raw"])
            self.assertIn("ccL75", manifest["unresolved"][0]["reason"])

    def test_stale_package_is_detected(self):
        with tempfile.TemporaryDirectory() as directory:
            repo = Path(directory)
            source = repo / "past" / "E05" / "source" / "raw_data" / "sample.xls"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"first")
            package = repo / "package"
            package.mkdir()
            (package / "manifest.json").write_text(json.dumps({"system": "E05", "sources": [{
                "path": "source/raw_data/sample.xls", "sha256": hashlib.sha256(b"first").hexdigest()}]}), encoding="utf-8")
            self.assertFalse(cleaner.validate_package_sources(repo, package))
            source.write_bytes(b"changed")
            self.assertEqual("source/raw_data/sample.xls", cleaner.validate_package_sources(repo, package)[0]["path"])


if __name__ == "__main__":
    unittest.main()
