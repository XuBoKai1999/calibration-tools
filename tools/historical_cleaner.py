"""Fail-closed prototype cleaner for representative E05/E07/E27 workbooks."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
import xlrd
from docx import Document


CONTEXT_FIELDS = ("scope_type", "scope_id", "key", "value", "unit", "note", "source")
RAW_FIELDS = {
    "E05": ("run_id", "point_id", "repeat_index", "setup_id", "date", "mode",
            "nominal_voltage", "nominal_voltage_unit", "observed_voltage",
            "observed_voltage_unit", "range_value", "range_unit", "resolution_value",
            "resolution_unit", "source"),
    "E07": ("run_id", "point_id", "repeat_index", "setup_id", "date", "mode",
            "nominal_primary_value", "nominal_primary_unit", "nominal_secondary_value",
            "nominal_secondary_unit", "excitation_percent", "observed_voltage",
            "observed_voltage_unit", "ratio_error", "ratio_error_unit", "phase_error",
            "phase_error_unit", "frequency_value", "frequency_unit", "burden_value",
            "burden_unit", "burden_power_factor", "range_value", "range_unit",
            "resolution_value", "resolution_unit", "source"),
    "E27": ("run_id", "point_id", "repeat_index", "setup_id", "date", "sample_role",
            "sample_id", "sheet_resistance", "sheet_resistance_unit", "source_stddev",
            "source_stddev_unit", "source_summary_note", "source"),
}

REPRESENTATIVES = {
    "E05": ("E170381.xls", "E230738.xls", "E250527.xls", "E260384.xlsx"),
    "E07": ("E150076.xls", "E200397.xls", "E250585.xls", "E260384.xlsx", "09807C02014.xls"),
    "E27": ("10307C02517-2543.xls", "E160049-2543.xls", "E230042.xls",
            "E260073.xlsx", "10407C02798-100.xls"),
}


class UnsupportedStructure(ValueError):
    pass


class AmbiguousStructure(UnsupportedStructure):
    pass


@dataclass
class Sheet:
    name: str
    rows: int
    cols: int
    get: object

    def cell(self, row: int, col: int):
        return self.get(row, col)


@dataclass
class Book:
    sheets: list[Sheet]
    datemode: int = 0


def load_book(path: Path) -> Book:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(str(path), formatting_info=True)
        sheets = []
        for source_sheet in book.sheets():
            def get(row, col, sheet=source_sheet):
                if row < 1 or col < 1 or row > sheet.nrows or col > sheet.ncols:
                    return None
                value = sheet.cell_value(row - 1, col - 1)
                return None if value == "" else value
            sheets.append(Sheet(source_sheet.name, source_sheet.nrows, source_sheet.ncols, get))
        return Book(sheets, book.datemode)
    if path.suffix.lower() == ".xlsx":
        book = openpyxl.load_workbook(path, data_only=False, read_only=False)
        sheets = []
        for source_sheet in book.worksheets:
            def get(row, col, sheet=source_sheet):
                value = sheet.cell(row, col).value
                return None if value == "" else value
            sheets.append(Sheet(source_sheet.title, source_sheet.max_row, source_sheet.max_column, get))
        return Book(sheets)
    raise UnsupportedStructure(f"unsupported workbook extension: {path.suffix}")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def tail(value) -> str:
    value = text(value)
    return value.split(":", 1)[-1].strip() if ":" in value else value


def source(raw_path: Path, sheet: Sheet, row: int, col: int) -> str:
    return f"raw:{raw_path.name}:{sheet.name}!{openpyxl.utils.get_column_letter(col)}{row}"


def iso_date(value, datemode=0) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and value > 20000:
        return xlrd.xldate_as_datetime(value, datemode).date().isoformat()
    match = re.search(r"(\d{2,4})[./-](\d{1,2})[./-](\d{1,2})", text(value))
    if not match:
        return ""
    year, month, day = map(int, match.groups())
    if year < 1911:
        year += 1911
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def parse_nominal(value, default_unit=""):
    match = re.search(r"([-+]?\d+(?:\.\d+)?)\s*([kKmM]?[vV])?", text(value))
    if not match:
        return None, ""
    return float(match.group(1)), match.group(2) or default_unit


def context_row(scope_type, scope_id, key, value, unit="", note="", locator=""):
    return dict(zip(CONTEXT_FIELDS, (scope_type, scope_id, key, text(value), unit, note, locator)))


def parse_environment(value, raw_path, sheet, row, col, scope_type="case", scope_id=""):
    values = re.findall(r"\d+(?:\.\d+)?", text(value))
    if len(values) < 3:
        return []
    locator = source(raw_path, sheet, row, col)
    rows = [context_row(scope_type, scope_id, "temperature", values[0], "C",
                        f"source tolerance/range: {values[1]}", locator)]
    humidity_index = 2 if len(values) >= 3 else 1
    note = f"source tolerance/range: {values[humidity_index + 1]}" if len(values) > humidity_index + 1 else ""
    rows.append(context_row(scope_type, scope_id, "humidity", values[humidity_index], "%RH", note, locator))
    return rows


def labelled_environment(raw_path, sheet, row, columns):
    rows = []
    for col in columns:
        value = text(sheet.cell(row, col))
        key = "temperature" if "溫度" in value else "humidity" if "濕度" in value else None
        if not key:
            continue
        values = re.findall(r"\d+(?:\.\d+)?", value)
        if not values:
            continue
        parsed = float(values[0])
        if key == "humidity" and not 0 <= parsed <= 100:
            raise AmbiguousStructure(f"impossible labelled humidity at {sheet.name}!{openpyxl.utils.get_column_letter(col)}{row}: {parsed}")
        rows.append(context_row("case", "", key, values[0], "C" if key == "temperature" else "%RH",
                                f"source tolerance/range: {values[1]}" if len(values) > 1 else "",
                                source(raw_path, sheet, row, col)))
    return rows


def workbook_strings(book: Book):
    for sheet in book.sheets:
        for row in range(1, sheet.rows + 1):
            for col in range(1, sheet.cols + 1):
                value = sheet.cell(row, col)
                if isinstance(value, str) and value.strip():
                    yield sheet, row, col, value.strip()


def main_sheet(book: Book) -> Sheet:
    nonempty = [sheet for sheet in book.sheets
                if any(sheet.cell(row, col) is not None
                       for row in range(1, sheet.rows + 1) for col in range(1, sheet.cols + 1))]
    if len(nonempty) != 1:
        raise UnsupportedStructure(f"expected one meaningful sheet, found {[sheet.name for sheet in nonempty]}")
    return nonempty[0]


def report_text(path: Path) -> str:
    document = Document(path)
    parts = [p.text for p in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def report_period(path: Path | None, system: str):
    if not path:
        return None
    matches = re.findall(r"(\d{4})年(\d{1,2})月(\d{1,2})日\D{0,5}(\d{4})年(\d{1,2})月(\d{1,2})日", report_text(path))
    if not matches:
        return None
    values = matches[0] if system == "E05" else matches[-1]
    return date(*map(int, values[:3])), date(*map(int, values[3:]))


def report_date_conflict(raw_dates, period, raw_path, report_path):
    dates = sorted(date.fromisoformat(value) for value in set(raw_dates) if value)
    if not dates or not period or (dates[0], dates[-1]) == period:
        return []
    return [{"concept": "calibration_dates", "raw": [value.isoformat() for value in dates],
             "report": [value.isoformat() for value in period],
             "raw_source": f"raw:{raw_path.name}:measurement-date headers",
             "report_source": f"report:{report_path.name}:calibration period"}]


def parse_trace(value):
    value = text(value)
    match = re.search(r"(E\d+[A-Z]?)\s*\(([^)]+)\)", value, re.I)
    return (match.group(1), iso_date(match.group(2))) if match else ("", "")


def add_instrument(rows, instrument_id, raw_path, sheet, role, instrument_type, serial_cell,
                   trace_cell=None, manufacturer="", model="", extra=()):
    serial_row, serial_col = serial_cell
    serial_text = text(sheet.cell(serial_row, serial_col))
    trace_text = text(sheet.cell(*trace_cell)) if trace_cell else ""
    if not serial_text and not trace_text and not model and not any(value not in (None, "") for _, value, _ in extra):
        return
    serial_match = re.search(r"S/N\s*:?\s*([A-Za-z0-9.-]+)", serial_text, re.I)
    if not serial_match:
        serial_match = re.search(r"(?:序號|：|:)\s*([A-Za-z0-9.-]+)\s*$", serial_text)
    serial = serial_match.group(1) if serial_match else serial_text if re.fullmatch(r"[A-Za-z0-9.-]+", serial_text) else ""
    locator = source(raw_path, sheet, serial_row, serial_col)
    properties = (("instrument_role", role, ""), ("instrument_type", instrument_type, ""),
                  ("manufacturer", manufacturer, ""), ("model", model, ""),
                  ("serial_number", serial, ""))
    for key, value, unit in properties + tuple(extra):
        if value not in (None, ""):
            rows.append(context_row("instrument", instrument_id, key, value, unit, "", locator))
    if trace_cell:
        trace_row, trace_col = trace_cell
        trace_id, trace_date = parse_trace(sheet.cell(trace_row, trace_col))
        trace_locator = source(raw_path, sheet, trace_row, trace_col)
        if trace_id:
            rows.append(context_row("instrument", instrument_id, "traceability_id", trace_id, locator=trace_locator))
        if trace_date:
            rows.append(context_row("instrument", instrument_id, "traceability_date", trace_date, locator=trace_locator))


def add_measurement_time(rows, raw_path, book, run_id="run_01"):
    seen = set()
    for sheet, row, col, value in workbook_strings(book):
        compact = value.replace(" ", "")
        if "~" in compact and re.search(r"(?:AM|PM|A\.M|P\.M).*\d{1,2}:\d{2}", compact, re.I) and value not in seen:
            rows.append(context_row("run", run_id, "measurement_time", value, note="verbatim",
                                    locator=source(raw_path, sheet, row, col)))
            seen.add(value)


def unresolved_labelled_metadata(raw_path, book):
    items = []
    for sheet, row, col, value in workbook_strings(book):
        label = next((label for label in ("廠家", "送校單位", "校正者", "製表者") if label in value), None)
        if not label:
            continue
        candidate = tail(value)
        if candidate == value or candidate == label:
            candidate = text(sheet.cell(row, col + 1))
        items.append({"severity": "warning", "concept": "unmapped_labelled_metadata", "label": label,
                      "value": candidate, "source": source(raw_path, sheet, row, col),
                      "reason": "no approved canonical key; preserved for human mapping"})
    return list({json.dumps(item, sort_keys=True): item for item in items}.values())


def instrument_identity(value: str):
    match = re.search(r"\b([A-Z][A-Z0-9.-]*)\s*[:/]\s*([A-Za-z0-9.-]+)\s*/\s*([A-Za-z0-9.-]+)", value)
    return match.groups() if match else None


def header_context(raw_path, sheet, positions):
    rows = [context_row("case", "", "work_order_id", raw_path.stem, locator=f"raw:{raw_path.name}:filename")]
    for key, row, col in positions:
        value = tail(sheet.cell(row, col))
        if value:
            rows.append(context_row("case", "", key, value, locator=source(raw_path, sheet, row, col)))
    return rows


def detect_generation(system: str, path: Path, book: Book) -> str:
    nonempty = [s for s in book.sheets if any(s.cell(r, c) is not None for r in range(1, s.rows + 1) for c in range(1, s.cols + 1))]
    if system == "E05" and len(nonempty) == 1:
        sheet = nonempty[0]
        if path.suffix.lower() == ".xlsx" and sheet.rows == 63 and 15 <= sheet.cols <= 16:
            return "E05-G4"
        if path.suffix.lower() == ".xls" and sheet.rows == 63 and sheet.cols == 24:
            return "E05-G3"
        if path.suffix.lower() == ".xls" and sheet.rows == 17 and 13 <= sheet.cols <= 20:
            return "E05-G2"
        if path.suffix.lower() == ".xls" and sheet.rows <= 24 and 10 <= sheet.cols <= 12:
            return "E05-G1"
    if system == "E07" and len(nonempty) == 1:
        sheet = nonempty[0]
        if path.suffix.lower() == ".xlsx" and sheet.rows == 63 and 18 <= sheet.cols <= 20:
            return "E07-G4"
        if path.suffix.lower() == ".xls" and sheet.rows == 63 and sheet.cols == 24:
            return "E07-G3"
        if path.suffix.lower() == ".xls" and sheet.rows <= 24 and sheet.cols >= 15 and "/" in text(sheet.cell(8, 1)):
            return "E07-G1"
        if path.suffix.lower() == ".xls" and sheet.rows <= 24 and 9 <= sheet.cols <= 12:
            return "E07-G2"
    if system == "E27":
        if len(nonempty) > 1:
            return "E27-MULTI"
        if len(nonempty) == 1:
            sheet = nonempty[0]
            if path.suffix.lower() == ".xlsx" and sheet.rows == 48 and sheet.cols == 24:
                return "E27-G3"
            if path.suffix.lower() == ".xls" and sheet.rows == 37 and sheet.cols in (9, 10):
                return "E27-G2" if not text(sheet.cell(3, 8)).endswith("V2023") else "E27-G3"
            if path.suffix.lower() == ".xls" and sheet.rows == 32 and sheet.cols == 9:
                return "E27-G1"
    raise UnsupportedStructure("workbook shape/labels do not match an approved representative generation")


def clean_e05(raw_path: Path, book: Book, generation: str, report_docx: Path | None):
    sheet = main_sheet(book)
    positions = (("dut_name", 5, 3), ("dut_manufacturer", 5, 4), ("dut_model", 5, 5),
                 ("dut_serial_number", 5, 6)) if generation == "E05-G4" else (
                    ("dut_name", 5, 2), ("dut_manufacturer", 5, 3),
                    ("dut_model", 5, 4), ("dut_serial_number", 5, 5))
    rows = header_context(raw_path, sheet, positions)
    env_row = 15 if generation == "E05-G4" else 23 if generation == "E05-G3" else 16
    rows.extend(parse_environment(sheet.cell(env_row, 1), raw_path, sheet, env_row, 1))
    mode = "divider_output" if generation == "E05-G4" else "direct_voltage"
    rows.append(context_row("run", "run_01", "mode", mode, locator=f"mapping:{generation}"))
    add_measurement_time(rows, raw_path, book)
    trace_meter_col = 8 if generation == "E05-G4" else 7
    add_instrument(rows, "standard_divider_01", raw_path, sheet, "standard_divider",
                   "high_voltage_divider", (6, 3), (6, 5), manufacturer="PARK")
    add_instrument(rows, "standard_voltmeter_01", raw_path, sheet, "standard_voltmeter",
                   "voltmeter", (6, 6), (6, trace_meter_col))
    raw = []
    reading_start_col = 2
    reading_end_col = 5
    observed_unit = "V" if mode == "divider_output" else "kV"
    for row in range(8, min(sheet.rows, 23) + 1):
        if not isinstance(sheet.cell(row, 1), str) or not re.search(r"\bkV\b", sheet.cell(row, 1), re.I):
            continue
        nominal, nominal_unit = parse_nominal(sheet.cell(row, 1), "kV")
        readings = [sheet.cell(row, col) for col in range(reading_start_col, reading_end_col + 1)]
        if nominal is None or not any(number(v) is not None for v in readings):
            continue
        resolution = sheet.cell(row, 8) if generation in ("E05-G2", "E05-G3") else None
        for repeat, col in enumerate(range(reading_start_col, reading_end_col + 1), 1):
            observed = number(sheet.cell(row, col))
            if observed is None:
                continue
            raw.append({"run_id": "run_01", "point_id": f"point_{row-7:02d}",
                        "repeat_index": repeat, "setup_id": "", "date": iso_date(sheet.cell(7, col), book.datemode),
                        "mode": mode, "nominal_voltage": nominal, "nominal_voltage_unit": nominal_unit,
                        "observed_voltage": observed, "observed_voltage_unit": observed_unit,
                        "range_value": "", "range_unit": "", "resolution_value": text(resolution),
                        "resolution_unit": "kV" if resolution not in (None, "") else "",
                        "source": source(raw_path, sheet, row, col)})

    raw_identities = []
    for item_sheet, row, col, value in workbook_strings(book):
        identity = instrument_identity(value)
        if identity and (value.startswith("*") or row > 15):
            raw_identities.append((identity, source(raw_path, item_sheet, row, col), value))
    report_identities = []
    if report_docx:
        for match in re.finditer(r"\b([A-Z][A-Z0-9.-]*)\s*[:/]\s*([A-Za-z0-9.-]+)\s*/\s*([A-Za-z0-9.-]+)", report_text(report_docx)):
            report_identities.append((match.groups(), f"report:{report_docx.name}:text", match.group(0)))
    chosen = raw_identities[0] if raw_identities else report_identities[0] if report_identities else None
    if chosen:
        identity, locator, original = chosen
        for key, value in zip(("manufacturer", "model", "serial_number"), identity):
            rows.append(context_row("instrument", "meter_01", key, value, locator=locator))
        role = "customer_meter" if chosen[2].startswith("*") else "auxiliary_meter"
        rows.append(context_row("instrument", "meter_01", "instrument_role", role,
                                note=f"verbatim: {original}", locator=locator))
    conflicts = report_date_conflict((row["date"] for row in raw), report_period(report_docx, "E05"), raw_path, report_docx)
    if raw_identities and report_identities and raw_identities[0][0] != report_identities[0][0]:
        conflicts.append({"concept": "customer_meter_identity", "raw": raw_identities[0][0],
                          "report": report_identities[0][0], "raw_source": raw_identities[0][1],
                          "report_source": report_identities[0][1]})
    if not raw:
        raise UnsupportedStructure("no E05 observations found in approved main table")
    return raw, rows, conflicts, []


def parse_burden(value):
    values = re.findall(r"\d+(?:\.\d+)?", text(value))
    if "VA" not in text(value).upper() or not values:
        return None
    return (values[-1], values[0] if len(values) > 1 else "", text(value))


def clean_e07(raw_path: Path, book: Book, generation: str, report_docx: Path | None):
    sheet = main_sheet(book)
    rows = header_context(raw_path, sheet, (("dut_name", 3, 2), ("dut_model", 3, 5),
                                            ("dut_serial_number", 3, 9)))
    rows.extend(labelled_environment(raw_path, sheet, 5, (1, 4, 8)))
    raw, unresolved = [], []
    transformer = generation == "E07-G1"
    mode = "voltage_transformer_ratio_phase" if transformer else "divider_output" if generation == "E07-G4" else "ac_high_voltage"
    rows.append(context_row("run", "run_01", "mode", mode, locator=f"mapping:{generation}"))
    add_measurement_time(rows, raw_path, book)
    if transformer:
        add_instrument(rows, "standard_transformer_01", raw_path, sheet, "standard_transformer",
                       "voltage_transformer", (4, 10), (4, 11))
    else:
        trace_col = 9 if generation in ("E07-G2", "E07-G4") else 8
        add_instrument(rows, "standard_transformer_01", raw_path, sheet, "standard_transformer",
                       "voltage_transformer", (4, 6), (4, trace_col))
        add_instrument(rows, "standard_ac_voltmeter_01", raw_path, sheet, "standard_ac_voltmeter",
                       "ac_voltmeter", (5, 6), (5, trace_col))
    period = report_period(report_docx, "E07")
    if not transformer:
        frequency_match = re.search(r"(\d+(?:\.\d+)?)\s*Hz", text(sheet.cell(7, 2)), re.I)
        if frequency_match:
            rows.append(context_row("run", "run_01", "frequency", frequency_match.group(1), "Hz",
                                    locator=source(raw_path, sheet, 7, 2)))
        for row in range(8, min(sheet.rows, 22) + 1):
            if not isinstance(sheet.cell(row, 1), str) or not re.search(r"\bkV\b", sheet.cell(row, 1), re.I):
                continue
            nominal, nominal_unit = parse_nominal(sheet.cell(row, 1), "kV")
            if nominal is None:
                continue
            for repeat, col in enumerate(range(3, 7), 1):
                observed = number(sheet.cell(row, col))
                if observed is None:
                    continue
                raw_date = iso_date(sheet.cell(6, col), book.datemode)
                if not raw_date and period:
                    partial = re.search(r"(\d{1,2})月(\d{1,2})日", text(sheet.cell(6, col)))
                    if partial:
                        recovered = date(period[0].year, int(partial.group(1)), int(partial.group(2)))
                        if period[0] <= recovered <= period[1]:
                            raw_date = recovered.isoformat()
                raw.append({"run_id": "run_01", "point_id": f"point_{row-7:02d}", "repeat_index": repeat,
                            "setup_id": "", "date": raw_date, "mode": mode,
                            "nominal_primary_value": nominal, "nominal_primary_unit": nominal_unit,
                            "nominal_secondary_value": "", "nominal_secondary_unit": "",
                            "excitation_percent": text(sheet.cell(row, 2)), "observed_voltage": observed,
                            "observed_voltage_unit": "V" if mode == "divider_output" else "kV",
                            "ratio_error": "", "ratio_error_unit": "", "phase_error": "",
                            "phase_error_unit": "", "frequency_value": "", "frequency_unit": "",
                            "burden_value": "", "burden_unit": "", "burden_power_factor": "",
                            "range_value": "", "range_unit": "", "resolution_value": "",
                            "resolution_unit": "", "source": source(raw_path, sheet, row, col)})
    else:
        footer_entry = next(((row, col, parse_burden(value)) for item_sheet, row, col, value in workbook_strings(book)
                             if item_sheet is sheet and parse_burden(value)
                             and not re.fullmatch(r"\s*\d+(?:\.\d+)?\s*(?:k?V)?\s*/\s*\d+(?:\.\d+)?\s*(?:k?V)?\s*",
                                                  text(sheet.cell(row, 1)), re.I)), None)
        table_has_burden = any(parse_burden(sheet.cell(row, 2)) and "/" in text(sheet.cell(row, 1))
                               for row in range(8, sheet.rows + 1))
        if footer_entry and not table_has_burden:
            burden_row, burden_col, footer_burden = footer_entry
            burden, power_factor, original = footer_burden
            rows.append(context_row("run", "run_01", "burden", burden, "VA", f"verbatim: {original}",
                                    source(raw_path, sheet, burden_row, burden_col)))
            rows.append(context_row("run", "run_01", "burden_power_factor", power_factor, "", "",
                                    source(raw_path, sheet, burden_row, burden_col)))
        setup_ids = {}
        for row in range(8, sheet.rows + 1):
            ratio = text(sheet.cell(row, 1))
            if not re.fullmatch(r"\s*\d+(?:\.\d+)?\s*(?:k?V)?\s*/\s*\d+(?:\.\d+)?\s*(?:k?V)?\s*", ratio, re.I):
                continue
            burden = parse_burden(sheet.cell(row, 2))
            key = (ratio, burden[0] if burden else "", burden[1] if burden else "")
            if key not in setup_ids:
                setup_id = f"setup_{len(setup_ids)+1:02d}"
                setup_ids[key] = setup_id
                rows.append(context_row("setup", setup_id, "ratio_text", ratio, locator=source(raw_path, sheet, row, 1)))
                if burden:
                    rows.append(context_row("setup", setup_id, "burden", burden[0], "VA", f"verbatim: {burden[2]}", source(raw_path, sheet, row, 2)))
                    rows.append(context_row("setup", setup_id, "burden_power_factor", burden[1], "", "", source(raw_path, sheet, row, 2)))
            setup_id = setup_ids[key]
            match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*(k?V)?\s*/\s*(\d+(?:\.\d+)?)\s*(k?V)?\s*", ratio, re.I)
            primary = secondary = primary_unit = secondary_unit = ""
            if match:
                primary, primary_unit, secondary, secondary_unit = match.groups()
            if not primary_unit or not secondary_unit:
                unresolved.append({"severity": "error", "concept": "ratio_units", "value": ratio,
                                   "source": source(raw_path, sheet, row, 1), "reason": "ratio units are not explicit"})
            pairs = ((4, 5), (6, 7), (8, 9), (10, 11))
            for repeat, (ratio_col, phase_col) in enumerate(pairs, 1):
                ratio_value, phase_value = number(sheet.cell(row, ratio_col)), number(sheet.cell(row, phase_col))
                if ratio_value is None and phase_value is None:
                    continue
                raw.append({"run_id": "run_01", "point_id": f"point_{row-7:02d}", "repeat_index": repeat,
                            "setup_id": setup_id, "date": iso_date(sheet.cell(6, ratio_col), book.datemode), "mode": mode,
                            "nominal_primary_value": primary, "nominal_primary_unit": primary_unit,
                            "nominal_secondary_value": secondary, "nominal_secondary_unit": secondary_unit,
                            "excitation_percent": text(sheet.cell(row, 3 if sheet.cell(row, 3) is not None else 2)),
                            "observed_voltage": "", "observed_voltage_unit": "", "ratio_error": text(ratio_value),
                            "ratio_error_unit": "%", "phase_error": text(phase_value), "phase_error_unit": "min",
                            "frequency_value": "", "frequency_unit": "", "burden_value": "", "burden_unit": "",
                            "burden_power_factor": "", "range_value": "", "range_unit": "",
                            "resolution_value": "", "resolution_unit": "", "source": source(raw_path, sheet, row, ratio_col)})
        unresolved.append({"severity": "warning", "concept": "frequency", "source": f"raw:{raw_path.name}",
                           "reason": "frequency is not explicit in the raw transformer table; no default inserted"})
    if not raw:
        raise UnsupportedStructure("no E07 observations found in approved main table")
    conflicts = report_date_conflict((row["date"] for row in raw), period, raw_path, report_docx)
    if generation == "E07-G4" and period and all(row["date"] for row in raw):
        rows.append(context_row("run", "run_01", "note", "Observation month/day headings completed with report year",
                                note=f"report period {period[0].isoformat()} through {period[1].isoformat()}",
                                locator=f"raw:{raw_path.name}:Sheet1!C6:F6; report:{report_docx.name}:calibration period"))
    unique_unresolved = list({json.dumps(item, sort_keys=True): item for item in unresolved}.values())
    return raw, rows, conflicts, unique_unresolved


def clean_e27(raw_path: Path, book: Book, generation: str):
    if generation == "E27-MULTI":
        names = [sheet.name for sheet in book.sheets]
        raise AmbiguousStructure(f"multi-sheet roles are not approved; sheets={names}")
    sheet = main_sheet(book)
    rows = header_context(raw_path, sheet, (("dut_name", 7, 2), ("dut_manufacturer", 8, 2),
                                            ("dut_model", 9, 2), ("dut_serial_number", 10, 2),
                                            ("material_type", 11, 2), ("thickness", 12, 2)))
    reference_id = text(sheet.cell(17, 7))
    if reference_id:
        rows.append(context_row("instrument", "reference_wafer_01", "reference_wafer_id", reference_id,
                                locator=source(raw_path, sheet, 17, 7)))
        rows.append(context_row("instrument", "reference_wafer_01", "instrument_role", "reference_wafer",
                                locator=source(raw_path, sheet, 17, 7)))
        rows.append(context_row("instrument", "reference_wafer_01", "model", reference_id,
                                locator=source(raw_path, sheet, 17, 7)))
    reference_serial = text(sheet.cell(17, 8))
    if reference_serial:
        rows.append(context_row("instrument", "reference_wafer_01", "serial_number", reference_serial,
                                locator=source(raw_path, sheet, 17, 8)))
    certificate_value = parse_nominal(sheet.cell(19, 7), "ohm")[0]
    certificate_uncertainty = parse_nominal(sheet.cell(19, 8), "ohm")[0]
    if certificate_value is not None:
        rows.append(context_row("instrument", "reference_wafer_01", "certificate_value", certificate_value, "ohm",
                                locator=source(raw_path, sheet, 19, 7)))
    if certificate_uncertainty is not None:
        rows.append(context_row("instrument", "reference_wafer_01", "certificate_uncertainty", certificate_uncertainty, "ohm",
                                locator=source(raw_path, sheet, 19, 8)))
    add_instrument(rows, "standard_meter_01", raw_path, sheet, "standard_meter", "voltmeter",
                   (9, 7), manufacturer="", model=text(sheet.cell(9, 6)))
    meter_trace = text(sheet.cell(9, 8))
    if meter_trace:
        rows.append(context_row("instrument", "standard_meter_01", "traceability_id", meter_trace,
                                locator=source(raw_path, sheet, 9, 8)))
    meter_date = iso_date(sheet.cell(9, 9), book.datemode)
    if meter_date:
        rows.append(context_row("instrument", "standard_meter_01", "traceability_date", meter_date,
                                locator=source(raw_path, sheet, 9, 9)))
    for index, resistor_row in enumerate(range(11, 16), 1):
        nominal, nominal_unit = parse_nominal(sheet.cell(resistor_row, 6), "ohm")
        add_instrument(rows, f"standard_resistor_{index:02d}", raw_path, sheet, "standard_resistor", "resistor",
                       (resistor_row, 7), extra=(("nominal_value", nominal, nominal_unit),))
        trace = text(sheet.cell(resistor_row, 8))
        if trace:
            rows.append(context_row("instrument", f"standard_resistor_{index:02d}", "traceability_id", trace,
                                    locator=source(raw_path, sheet, resistor_row, 8)))
        trace_date = iso_date(sheet.cell(resistor_row, 9), book.datemode)
        if trace_date:
            rows.append(context_row("instrument", f"standard_resistor_{index:02d}", "traceability_date", trace_date,
                                    locator=source(raw_path, sheet, resistor_row, 9)))
    probe = text(sheet.cell(20, 2))
    if probe:
        rows.append(context_row("instrument", "probe_01", "instrument_role", "probe", locator=source(raw_path, sheet, 20, 2)))
        rows.append(context_row("instrument", "probe_01", "serial_number", probe, locator=source(raw_path, sheet, 20, 2)))
    raw, run_for_date, counts = [], {}, defaultdict(int)
    carried_date = ""
    for row in range(23, min(sheet.rows, 26) + 1):
        parsed = iso_date(sheet.cell(row, 1), book.datemode)
        if parsed:
            carried_date = parsed
        if not carried_date:
            continue
        run_id = run_for_date.setdefault(carried_date, f"run_{len(run_for_date)+1:02d}")
        for role, value_col, std_col, sample_id in (("dut", 2, 3, text(sheet.cell(10, 2)) or raw_path.stem),
                                                     ("reference_wafer", 5, 6, reference_id)):
            value = number(sheet.cell(row, value_col))
            if value is None:
                continue
            counts[(run_id, role)] += 1
            raw.append({"run_id": run_id, "point_id": f"{role}_01", "repeat_index": counts[(run_id, role)],
                        "setup_id": "setup_01", "date": carried_date, "sample_role": role,
                        "sample_id": sample_id, "sheet_resistance": value, "sheet_resistance_unit": "ohm/square",
                        "source_stddev": text(sheet.cell(row, std_col)),
                        "source_stddev_unit": "1" if sheet.cell(row, std_col) is not None else "",
                        "source_summary_note": "source-provided relative standard deviation (source displayed as percent); underlying individual readings are absent" if sheet.cell(row, std_col) is not None else "",
                        "source": source(raw_path, sheet, row, value_col)})
    for index, (run_date, run_id) in enumerate(run_for_date.items(), 5):
        if index <= 6:
            if number(sheet.cell(index, 7)) is not None:
                rows.append(context_row("run", run_id, "temperature", sheet.cell(index, 7), "C", "", source(raw_path, sheet, index, 7)))
            if number(sheet.cell(index, 8)) is not None:
                rows.append(context_row("run", run_id, "humidity", sheet.cell(index, 8), "%RH", "", source(raw_path, sheet, index, 8)))
    rows.append(context_row("setup", "setup_01", "reference_wafer_id", reference_id, locator=source(raw_path, sheet, 17, 7)))
    rows.append(context_row("setup", "setup_01", "instrument_id", "probe_01", locator=source(raw_path, sheet, 20, 2)))
    unresolved = []
    if re.search(r"-\d+$", raw_path.stem):
        unresolved.append({"severity": "warning", "concept": "filename_suffix", "value": raw_path.stem.rsplit("-", 1)[1],
                           "source": f"raw:{raw_path.name}:filename", "reason": "suffix semantics are not approved"})
    if not raw:
        raise UnsupportedStructure("no E27 observations found in approved measurement rows")
    return raw, rows, [], unresolved


def write_csv(path: Path, fieldnames, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def find_reports(system_root: Path, reference_id: str):
    report_root = system_root / "source" / "report"
    return sorted((p for p in report_root.glob(f"{reference_id}.*") if p.suffix.lower() in (".doc", ".docx")),
                  key=lambda p: (p.suffix.lower() != ".docx", p.name.lower()))


def convert_one(repo_root: Path, system: str, raw_path: Path, output_root: Path):
    reference_id = raw_path.stem
    destination = output_root / system / reference_id
    if destination.exists():
        raise FileExistsError(f"destination already exists: {destination}")
    system_root = repo_root / "past" / system
    reports = find_reports(system_root, reference_id)
    docx = next((p for p in reports if p.suffix.lower() == ".docx"), None)
    sources = [{"path": raw_path.relative_to(system_root).as_posix(), "type": "raw", "sha256": sha256(raw_path)}]
    sources.extend({"path": p.relative_to(system_root).as_posix(), "type": "report", "sha256": sha256(p)} for p in reports)
    destination.mkdir(parents=True)
    artifacts = {"raw": None, "context": None, "report": None}
    conflicts, unresolved, generation = [], [], "unsupported"
    status = "FAIL_UNSUPPORTED"
    try:
        book = load_book(raw_path)
        generation = detect_generation(system, raw_path, book)
        sources[0]["generation"] = generation
        if system == "E05":
            raw, context, conflicts, unresolved = clean_e05(raw_path, book, generation, docx)
        elif system == "E07":
            raw, context, conflicts, unresolved = clean_e07(raw_path, book, generation, docx)
        else:
            raw, context, conflicts, unresolved = clean_e27(raw_path, book, generation)
        unresolved.extend(unresolved_labelled_metadata(raw_path, book))
        if system == "E05" and reference_id == "E170381" and reports:
            conflicts.append({"concept": "known_raw_report_mismatch",
                              "raw": {"dates": sorted({row["date"] for row in raw}), "one_kv_readings": [row["observed_voltage"] for row in raw if row["point_id"] == "point_01"]},
                              "report": {"dates": ["2017-07-04", "2017-07-07"], "one_kv_result_approx": 0.998},
                              "raw_source": f"raw:{raw_path.name}:Sheet1!A7:E8",
                              "report_source": f"report:{reports[0].name}:human-reviewed calibration period/result",
                              "note": "legacy .doc is not automatically parsed; mismatch was confirmed during representative human review"})
        for row in context:
            if row["key"] == "humidity" and not 0 <= float(row["value"]) <= 100:
                raise AmbiguousStructure(f"humidity outside physical range: {row['value']}")
            if row["key"] in {"dut_name", "dut_manufacturer", "dut_model", "dut_serial_number"} and re.fullmatch(r"E\d+", row["value"], re.I):
                raise AmbiguousStructure(f"DUT identity column shift suspected: {row['key']}={row['value']}")
        current_hash = sha256(raw_path)
        if current_hash != sources[0]["sha256"]:
            raise AmbiguousStructure(f"source changed during conversion: expected {sources[0]['sha256']}, found {current_hash}")
        write_csv(destination / "raw.csv", RAW_FIELDS[system], raw)
        artifacts["raw"] = "raw.csv"
        if context:
            write_csv(destination / "context.csv", CONTEXT_FIELDS, context)
            artifacts["context"] = "context.csv"
        if docx:
            shutil.copyfile(docx, destination / "report.docx")
            artifacts["report"] = "report.docx"
        elif any(p.suffix.lower() == ".doc" for p in reports):
            unresolved.append({"severity": "warning", "concept": "legacy_report_conversion", "source": f"report:{reports[0].name}",
                               "reason": "only legacy .doc exists; canonical report was not invented"})
        if conflicts or any(item.get("severity") == "error" for item in unresolved):
            status = "FAIL_AMBIGUOUS"
        elif unresolved:
            status = "PASS_WITH_UNRESOLVED"
        else:
            status = "PASS"
    except AmbiguousStructure as exc:
        status = "FAIL_AMBIGUOUS"
        unresolved.append({"severity": "error", "concept": "ambiguous_structure", "source": f"raw:{raw_path.name}", "reason": str(exc)})
    except UnsupportedStructure as exc:
        unresolved.append({"severity": "error", "concept": "unsupported_structure", "source": f"raw:{raw_path.name}", "reason": str(exc)})
    manifest = {"schema_version": 1, "system": system, "reference_id": reference_id, "sources": sources,
                "artifacts": artifacts, "conflicts": conflicts, "unresolved": unresolved}
    (destination / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"reference_id": reference_id, "detected_generation": generation,
            "source_files": ";".join(item["path"] for item in sources),
            "raw_row_count": len(raw) if artifacts["raw"] else 0,
            "context_row_count": len(context) if artifacts["context"] else 0,
            "report_status": "COPIED_DOCX" if artifacts["report"] else "ABSENT",
            "conflict_count": len(conflicts), "unresolved_count": len(unresolved), "conversion_status": status}


def validate_package_sources(repo_root: Path, package: Path):
    manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
    system_root = repo_root / "past" / manifest["system"]
    stale = []
    for item in manifest["sources"]:
        path = system_root / item["path"]
        actual = sha256(path) if path.is_file() else None
        if actual != item["sha256"]:
            stale.append({"path": item["path"], "recorded": item["sha256"], "actual": actual})
    return stale


def run_representatives(repo_root: Path, output_root: Path):
    results = []
    for system, names in REPRESENTATIVES.items():
        for name in names:
            raw_path = repo_root / "past" / system / "source" / "raw_data" / name
            results.append(convert_one(repo_root, system, raw_path, output_root))
    write_csv(output_root / "validation.csv", tuple(results[0]), results)
    return results


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path)
    parser.add_argument("--representatives", action="store_true")
    parser.add_argument("--system", choices=tuple(RAW_FIELDS))
    parser.add_argument("--raw", type=Path)
    parser.add_argument("--validate-package", type=Path)
    args = parser.parse_args(argv)
    repo_root = Path(__file__).resolve().parents[1]
    if args.validate_package:
        stale = validate_package_sources(repo_root, args.validate_package.resolve())
        print(json.dumps({"status": "STALE" if stale else "CURRENT", "mismatches": stale}, ensure_ascii=False, indent=2))
        return 1 if stale else 0
    if not args.output_root:
        parser.error("--output-root is required for conversion")
    if args.representatives:
        results = run_representatives(repo_root, args.output_root.resolve())
    elif args.system and args.raw:
        results = [convert_one(repo_root, args.system, args.raw.resolve(), args.output_root.resolve())]
    else:
        parser.error("use --representatives or both --system and --raw")
    for result in results:
        print(", ".join(f"{key}={value}" for key, value in result.items()))
    return 1 if any(result["conversion_status"].startswith("FAIL") for result in results) else 0


if __name__ == "__main__":
    sys.exit(main())
